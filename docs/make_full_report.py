from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_font(run, size=10, bold=False, color=None):
    run.font.name = '맑은 고딕'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '맑은 고딕')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color

def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    font_size = 18 - (level * 2)
    set_font(run, size=font_size, bold=True, color=RGBColor(0, 0, 0))
    return h

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run, size=10, bold=bold)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    set_font(run, size=10)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))

def create_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    
    # 헤더 설정
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        run = hdr_cells[i].paragraphs[0].runs[0]
        set_font(run, size=9, bold=True, color=RGBColor(255, 255, 255))
        hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        shading_elm = hdr_cells[i]._element.get_or_add_tcPr()
        from docx.oxml.shared import OxmlElement
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), '4472C4') # 파란색 헤더
        shading_elm.append(shd)

    # 데이터 입력
    for row_data in data:
        row_cells = table.add_row().cells
        for i, item in enumerate(row_data):
            cell = row_cells[i]
            cell.text = str(item)
            if cell.paragraphs:
                run = cell.paragraphs[0].runs[0]
                set_font(run, size=9)
                # 중앙 정렬 (첫 열과 숫자 열 등)
                if i == 0 or (isinstance(item, (int, float)) and i > 0):
                     cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 컬럼 너비 설정
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Cm(width)

    doc.add_paragraph() # Spacer

def generate_full_report():
    doc = Document()
    
    # --- 타이틀 페이지 ---
    doc.add_paragraph('\n\n\n\n')
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('성층권 고가용성 센서 플랫폼 개발\n상세 기술 보고서')
    set_font(run, size=24, bold=True)
    
    doc.add_paragraph('\n')
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('STM32 Stratospheric Balloon Project\nTechnical Report Rev 4.2')
    set_font(run, size=14, color=RGBColor(80, 80, 80))
    
    doc.add_paragraph('\n\n\n\n\n\n')
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run('문서 번호: SB-REP-2026-001\n날짜: 2026년 1월 9일\n작성자: Antigravity AI')
    set_font(run, size=11)
    
    doc.add_page_break()

    # --- 1. 서론 ---
    add_heading(doc, '1. 서론 (Introduction)', level=1)
    add_heading(doc, '1.1 배경 및 목적', level=2)
    add_paragraph(doc, '본 프로젝트는 영하 60°C의 성층권(고도 30km) 환경에서 3시간 동안 비행하며 대기 데이터를 수집하는 고가용성 센서 플랫폼을 개발하는 것을 목표로 합니다. 극한 환경에서의 배터리 전압 강하, 센서 데이터 드리프트, 우주 방사선(SEU)으로 인한 시스템 락업 등의 결함을 극복하기 위해 STM32G431 기반의 내결함성(Fault-Tolerant) 아키텍처를 설계 및 구현하였습니다.')
    
    add_heading(doc, '1.2 미션 프로파일', level=2)
    add_bullet(doc, '운용 고도: 30km (성층권)')
    add_bullet(doc, '비행 시간: 약 3시간 (상승 1.5h + 하강 1.5h)')
    add_bullet(doc, '환경 조건: 최저 기온 -60°C, 기압 10mbar (0.01 atm)')

    # --- 2. 하드웨어 시스템 구성 ---
    add_heading(doc, '2. 하드웨어 시스템 구성', level=1)
    add_paragraph(doc, '시스템은 단일 고장점(SPOF)을 제거하고 노이즈 간섭을 최소화하기 위해 버스 분리 토폴로지를 적용하였으며, Upside와 Downside 보드로 물리적으로 구분됩니다.')

    add_heading(doc, '2.1 보드별 상세 구성', level=2)
    
    add_paragraph(doc, '1) Upside Board (환경 센서 및 메인 제어)', bold=True)
    upside_data = [
        ['구분', '부품명', '기능', '통신 인터페이스'],
        ['MCU', 'STM32G431CBU6', 'Main Controller (170MHz, Cortex-M4F)', '-'],
        ['Sensor', 'MS5611', '정밀 기압/고도 (10-1200mbar)', 'I2C3'],
        ['Sensor', 'SHT31-D', '온도/습도', 'I2C3'],
        ['Sensor', 'CM1107N', 'CO2 농도 (NDIR 방식)', 'I2C3'],
        ['Sensor', 'MCP9600', 'K-type 열전대 인터페이스', 'I2C3'],
        ['Sensor', 'SEN0321', '오존(O3) 가스 센서', 'I2C3'],
        ['Power', 'LDO', '3.3V 정전압 레귤레이터', '-'],
        ['Thermal', 'Heater Driver', 'PWM MOSFET 제어 (필름 히터)', 'TIM PWM']
    ]
    create_table(doc, upside_data[0], upside_data[1:], [3, 4, 6, 3])

    add_paragraph(doc, '\n2) Downside Board (비행 역학 및 안전)', bold=True)
    downside_data = [
        ['구분', '부품명', '기능', '통신 인터페이스'],
        ['Sensor', 'LSM6DSV16X', '6축 IMU (가속도+자이로)', 'I2C1'],
        ['Sensor', 'MLX90393', '3축 지자기 (Compass)', 'I2C1'],
        ['Sensor', 'GDK101', '감마선 방사선 센서', 'I2C1'],
        ['Safety', 'P-MOS Switch', 'Latch-up 복구용 전원 스위치', 'GPIO Output']
    ]
    create_table(doc, downside_data[0], downside_data[1:], [3, 4, 6, 3])

    add_paragraph(doc, '\n3) 독립 인터페이스 장치', bold=True)
    uart_data = [
        ['장치명', '모델', '기능', '연결'],
        ['GPS', 'XA1110', '위치 추적 및 1PPS 시간 동기화', 'UART1'],
        ['LoRa', 'RFM95W/SX1276', '장거리 텔레메트리 (915MHz)', 'UART2 + SPI'],
        ['Dust', 'PMS3003', '미세먼지 측정', 'UART3']
    ]
    create_table(doc, uart_data[0], uart_data[1:], [3, 4, 6, 3])

    # --- 3. FMEA 상세 분석 ---
    doc.add_page_break()
    add_heading(doc, '3. FMEA (고장 모드 및 영향 분석) 상세', level=1)
    add_paragraph(doc, '모든 서브시스템에 대해 식별된 잠재적 고장 모드와 이에 대한 위험도(RPN) 평가 결과입니다. RPN 20 이상은 중점 관리 항목으로 분류됩니다.')

    # FMEA 테이블 헤더
    fmea_headers = ['ID', '고장 모드', '원인', 'S', 'O', 'D', 'RPN', '완화 조치']
    col_widths = [1.5, 3, 5, 1, 1, 1, 1.5, 5]

    add_heading(doc, '3.1 센서 서브시스템 (S-Series)', level=2)
    sensor_fmea = [
        ['S-01', 'IMU 드리프트', '진동/온도 변화로 인한 바이어스 오차', 3, 3, 3, 27, 'Kalman Filter, 온도 보상'],
        ['S-02', 'IMU 락업', 'SEU(방사선)로 인한 Latch-up', 3, 2, 3, 18, 'WDT, P-MOS 전원 사이클'],
        ['S-03', 'IMU 노이즈', '기계적 진동으로 인한 데이터 오염', 2, 4, 2, 16, 'LPF, 기구적 방진 댐퍼'],
        ['S-04', 'GPS No Fix', '안테나 지향성 상실, 고도 제한', 4, 3, 2, 24, 'Barometer 백업 고도 사용'],
        ['S-05', 'GPS 데이터 에러', 'UART 통신 노이즈, 비트 반전', 2, 3, 1, 6, 'NMEA Checksum 검증'],
        ['S-06', 'GPS 파싱 실패', '버퍼 오버플로우', 2, 2, 2, 8, 'Circular Buffer, DMA 수신'],
        ['S-07', 'Baro 스파이크', '기류 불안정으로 인한 튀는 값', 2, 3, 2, 12, '이동 평균 필터, Validation'],
        ['S-08', 'Baro 오프셋', '급격한 온도 변화로 인한 영점 이동', 2, 2, 2, 8, 'GPS 고도와 상호 검증'],
        ['S-09', 'Baro 멈춤', 'I2C 버스 에러', 3, 1, 3, 9, 'FDIR 타임아웃 감지 후 리셋'],
        ['S-10', 'SHT31 동결', '성층권 결로로 인한 센서 마비', 3, 5, 1, 15, '내장 히터 활성화, 단열'],
        ['S-11', 'SHT31 결로', '하강 시 급격한 온도 변화', 2, 4, 2, 16, '하강 전 히터 가동 프로파일'],
        ['S-12', 'GDK101 노이즈', '전원 노이즈 유입', 2, 4, 2, 16, '이동 평균, 전원 필터'],
        ['S-13', 'GDK101 고장', '물리적 충격', 2, 1, 3, 6, '값 0 지속 시 FDIR 경고']
    ]
    create_table(doc, fmea_headers, sensor_fmea, col_widths)

    add_heading(doc, '3.2 전력 서브시스템 (P-Series)', level=2)
    power_fmea = [
        ['P-01', '배터리 저온', '-60°C 저항 급증(Voltage Sag)', 4, 4, 2, 32, '필름 히터(PID), 단열재 (Top Risk)'],
        ['P-02', '과방전', '미션 시간 초과, 전력 소비 과다', 5, 1, 5, 25, '저전압 부하 차단 로직'],
        ['P-03', '셀 불균형', '충전 불량으로 인한 용량 감소', 3, 2, 3, 18, '발사 전 전압 밸런싱 체크'],
        ['P-04', 'LDO 과열', '부하 급증으로 인한 열적 차단', 4, 2, 2, 16, 'PCB 방열 설계, 입력 분산'],
        ['P-05', 'LDO 리플', '입력 전압 변동 노이즈', 3, 3, 2, 18, '입출력 탄탈 커패시터 보강']
    ]
    create_table(doc, fmea_headers, power_fmea, col_widths)

    add_heading(doc, '3.3 통신 서브시스템 (C-Series)', level=2)
    comm_fmea = [
        ['C-01', 'LoRa 링크 손실', '거리 증가, 지형, 안테나 틀어짐', 4, 3, 2, 24, 'SD카드 로깅(Blackbox), 재전송'],
        ['C-02', '프레임 손상', '전송 중 노이즈로 인한 비트 반전', 2, 3, 2, 12, 'CRC16 체크섬 검증'],
        ['C-03', '주파수 드리프트', '저온 시 오실레이터 주파수 변이', 4, 2, 2, 16, 'TCXO 사용, 대역폭 여유 설정']
    ]
    create_table(doc, fmea_headers, comm_fmea, col_widths)

    add_heading(doc, '3.4 열제어 서브시스템 (T-Series)', level=2)
    thermal_fmea = [
        ['T-01', '히터 단선', '충격/진동으로 인한 물리적 파손', 4, 1, 5, 20, '커넥터 글루건 고정, 회로 이중화'],
        ['T-02', '과열 (Runaway)', '센서 오류, 제어 로직 실패', 3, 1, 5, 15, 'SW 안전 차단(80°C Limit)'],
        ['T-03', 'PWM 고장', '타이머 설정 오류', 3, 1, 3, 9, '타이머 상태 모니터링'],
        ['T-04', '단열재 파손', '외부 충격으로 인한 크랙', 3, 1, 4, 12, '테이핑 보강, 충격 흡수재']
    ]
    create_table(doc, fmea_headers, thermal_fmea, col_widths)

    add_heading(doc, '3.5 소프트웨어 (W-Series)', level=2)
    sw_fmea = [
        ['W-01', '스택 오버플로', '재귀 호출, 과도한 지역 변수', 5, 1, 4, 20, '정적 분석, 스택 모니터링'],
        ['W-02', '무한 루프', 'I2C 타임아웃 미처리', 4, 2, 2, 16, 'Watchdog Timer (WDT) 활성화'],
        ['W-03', '메모리 누수', '동적 할당 미해제', 4, 1, 4, 16, '정적 할당(No Malloc) 원칙'],
        ['W-04', 'ISR 지연', '인터럽트 과부하', 3, 3, 3, 27, '실행 시간 계측, 우선순위 최적화'],
        ['W-05', 'Kalman 발산', '비정상 입력값, 행렬 연산 오류', 3, 2, 2, 12, '입력값 검증, 리셋 로직'],
        ['W-06', 'Float 오류', 'NaN, Inf 발생', 3, 1, 2, 6, 'isnan() 검사 후 처리']
    ]
    create_table(doc, fmea_headers, sw_fmea, col_widths)

    # --- 4. FDIR 아키텍처 ---
    doc.add_page_break()
    add_heading(doc, '4. FDIR (Fault Detection, Isolation, Recovery)', level=1)
    
    add_heading(doc, '4.1 감지 및 복구 전략', level=2)
    add_paragraph(doc, 'FDIR 시스템은 1Hz 주기로 모든 센서의 상태를 점검하며, 3가지 감지 메커니즘(Timeout, Range, Continuity)과 4단계 복구 전략을 사용합니다.')
    
    add_paragraph(doc, '복구 레벨 상세:', bold=True)
    recovery_levels = [
        ['레벨', '동작', '설명'],
        ['L1', 'Retry', '단순 재시도 (일시적 통신 오류)'],
        ['L2', 'Soft Reset', '드라이버 재초기화 및 I2C 9-Clock Pulse (버스 해제)'],
        ['L3', 'Hard Reset', 'P-MOS 전원 사이클 (Latch-up 물리적 해제)'],
        ['L4', 'Isolation', '5회 이상 복구 실패 시 영구 비활성화 및 대체 값 사용']
    ]
    create_table(doc, recovery_levels[0], recovery_levels[1:], [2, 3, 11])

    add_heading(doc, '4.2 센서별 리셋 핀 매핑 (Hardware Reset)', level=2)
    add_paragraph(doc, 'L3 복구(Hard Reset)를 위해 각 센서에 할당된 P-MOS 제어 핀 정보입니다.')
    reset_pins = [
        ['센서', '인터페이스', '리셋 제어 핀', 'Timeout 설정'],
        ['LSM6DSV16X (IMU)', 'I2C1', 'PB13', '100 ms'],
        ['MLX90393 (Mag)', 'I2C1', 'PB14', '500 ms'],
        ['GDK101 (Rad)', 'I2C1', 'PB2', '3000 ms'],
        ['MS5611 (Baro)', 'I2C3', 'PA5', '1000 ms'],
        ['SHT31 (Env)', 'I2C3', 'PB11', '3000 ms'],
        ['MCP9600 (Temp)', 'I2C3', 'PA4', '3000 ms'],
        ['CM1107N (CO2)', 'I2C3', 'PB0', '5000 ms'],
        ['GPS (XA1110)', 'UART1', 'PA9 (RST)', '5000 ms']
    ]
    create_table(doc, reset_pins[0], reset_pins[1:], [4, 3, 4, 4])

    # --- 5. 구현 현황 ---
    doc.add_page_break()
    add_heading(doc, '5. 구현 및 검증 현황 (Implementation Status)', level=1)
    add_paragraph(doc, '2026-01-09 기준, Priority 0(Critical) 및 Priority 1(Stability) 항목의 구현 및 단위 테스트가 완료되었습니다.')

    add_heading(doc, '5.1 구현 상세 (Priority 0 & 1)', level=2)
    impl_data = [
        ['우선순위', '항목', '관련 파일', 'LOC', '상세 내용'],
        ['P0', 'Sensors_Reset()', 'sensors.c', '110', '9개 센서 L2/L3 통합 복구 로직 구현'],
        ['P0', 'I2C Bus Recovery', 'bsp_i2c.c', '104', 'GPIO 비트뱅잉을 통한 SCL 9-Clock 생성 및 버스 해제'],
        ['P0', 'Telemetry Send', 'telemetry.c', '-', '132byte 바이너리 패킷 구성 및 UART DMA 전송'],
        ['P1', 'GPS NMEA CRC', 'xa1110.c', '45', 'XOR Checksum 계산 및 검증 (minmea_check)'],
        ['P1', '저전압 보호', 'app_power.c', '50', '2.7V Cutoff / 2.9V Recovery 히스테리시스 제어'],
        ['P1', '1PPS 동기화', 'pps_capture.c', '100', 'GPS 1PPS 인터럽트 기반 50Hz 슬롯 타이밍 보정']
    ]
    create_table(doc, impl_data[0], impl_data[1:], [2, 4, 3, 2, 8])

    add_heading(doc, '5.2 검증 결과 (Unit Test)', level=2)
    add_paragraph(doc, '총 27개의 단위 테스트 케이스를 100% 통과하였습니다.')
    test_data = [
        ['모듈', '테스트 항목', 'Pass/Fail'],
        ['FDIR', 'Timeout 감지 및 상태 전이 확인', 'PASS (4/4)'],
        ['Kalman', '행렬 연산 및 수렴성 확인', 'PASS (4/4)'],
        ['Control', 'PID 온도 제어 응답 확인', 'PASS (5/5)'],
        ['Power', '저전압 히스테리시스 동작 확인', 'PASS (5/5)'],
        ['Protocol', 'GPS NMEA 파싱 및 CRC 검증', 'PASS (9/9)']
    ]
    create_table(doc, test_data[0], test_data[1:], [4, 8, 3])

    # --- 6. 향후 계획 ---
    doc.add_page_break()
    add_heading(doc, '6. 향후 계획 (Next Steps)', level=1)
    
    add_heading(doc, '6.1 하드웨어 통합 테스트 (Priority 1.5)', level=2)
    add_paragraph(doc, '소프트웨어 구현이 완료된 기능을 실제 하드웨어에서 검증하는 단계입니다. (기간: 1주일)')
    add_bullet(doc, 'P-MOS 전원 사이클 동작 시 Inrush Current 및 전압 강하 측정')
    add_bullet(doc, 'GPIO 제어 신호와 센서 응답 간의 타이밍 마진 측정 (오실로스코프)')
    add_bullet(doc, 'I2C 강제 Short 시 복구 파형 분석')

    add_heading(doc, '6.2 시스템 검증 (Priority 2)', level=2)
    add_paragraph(doc, '가상의 비행 환경을 모사하여 시스템 안정성을 검증합니다. (기간: 3-5일)')
    add_bullet(doc, 'HITL (Hardware-In-The-Loop): ESP32 Mock System을 이용해 센서 고장, 통신 단절 등을 강제로 주입')
    add_bullet(doc, 'SITL (Software-In-The-Loop): RS41 실제 비행 데이터(5.4km, -50°C)를 주입하여 알고리즘 검증')

    add_heading(doc, '6.3 최종 승인 및 발사', level=2)
    add_bullet(doc, '환경 챔버 테스트 (-40°C ~ +60°C 열충격 시험)')
    add_bullet(doc, 'FIT (Flight Integration Test): 배터리 완충 상태에서 3시간 연속 구동')
    add_bullet(doc, 'FM (Flight Model) 전환 및 최종 발사 승인')

    # 저장
    file_name = 'STM32_Final_Report_Complete.docx'
    doc.save(file_name)
    print(f"'{file_name}' 파일 생성이 완료되었습니다.")

if __name__ == '__main__':
    generate_full_report()